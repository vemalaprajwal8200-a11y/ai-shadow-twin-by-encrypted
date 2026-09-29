"""Turn uploaded PPTX slides and PDF pages into ordered course items."""

import json
import sys
from pathlib import Path
from uuid import uuid4

from pptx import Presentation
from pypdf import PdfReader

from app.models import Item


def _item(course_id: str, order: int, text: str) -> Item:
    return Item(
        itemId=str(uuid4()),
        courseId=course_id,
        orderIndex=order,
        type="slide",
        text=text,
    )


def parse_pptx(path: str | Path, course_id: str) -> list[Item]:
    items = []
    for order, slide in enumerate(Presentation(path).slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if shape.has_table:
                parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        text = "\n".join(part.strip() for part in parts if part.strip())
        if text:
            items.append(_item(course_id, order, text))
    return items


def parse_pdf(path: str | Path, course_id: str) -> list[Item]:
    items = []
    for order, page in enumerate(PdfReader(path).pages, start=1):
        text = (page.extract_text() or "").strip()
        items.append(_item(course_id, order, text))
    return items


if __name__ == "__main__":
    file_path = Path(sys.argv[1])
    course_id = sys.argv[2] if len(sys.argv) > 2 else "local-test"
    parser = parse_pptx if file_path.suffix.lower() == ".pptx" else parse_pdf
    print(json.dumps([item.model_dump() for item in parser(file_path, course_id)], indent=2))